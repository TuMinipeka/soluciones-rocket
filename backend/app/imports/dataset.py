"""Validated, repeatable import of authorized room and equipment datasets."""

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pydantic import ValidationError as PydanticValidationError
from sqlalchemy.orm import Session

from app.modules.resources.models import Resource
from app.modules.resources.schemas import ResourceCreate
from app.modules.spaces.models import Space
from app.services.availability import OFFICIAL_SPACES


FIELDS = {
    "resources": ({"id", "name", "category"}, {"condition"}),
    "spaces": ({"id", "name", "capacity"}, set()),
}


def _read_rows(path: Path) -> list[tuple[int, dict[str, Any]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            reader = csv.DictReader(source)
            if not reader.fieldnames:
                raise ValueError("El CSV no contiene encabezados.")
            return [(number, row) for number, row in enumerate(reader, start=2)]
    if suffix == ".json":
        with path.open("r", encoding="utf-8-sig") as source:
            payload = json.load(source)
        if isinstance(payload, dict):
            payload = payload.get("records")
        if not isinstance(payload, list):
            raise ValueError("El JSON debe ser una lista o contener una lista 'records'.")
        return [(number, row) for number, row in enumerate(payload, start=1)]
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            values = workbook.active.iter_rows(values_only=True)
            headers = next(values, None)
            if not headers:
                raise ValueError("El Excel no contiene encabezados.")
            names = [str(value).strip() if value is not None else "" for value in headers]
            return [
                (number, dict(zip(names, cells)))
                for number, cells in enumerate(values, start=2)
                if any(value is not None for value in cells)
            ]
        finally:
            workbook.close()
    raise ValueError("Formato no admitido. Usa .csv, .json o .xlsx.")


def _clean(row: dict[str, Any], kind: str) -> dict[str, Any]:
    if not isinstance(row, dict):
        raise ValueError("La fila debe ser un objeto con columnas.")
    if any(not isinstance(key, str) for key in row):
        raise ValueError("La fila contiene un encabezado vacío o inválido.")
    required, optional = FIELDS[kind]
    names = set(row)
    if not required <= names:
        raise ValueError(f"Faltan columnas: {', '.join(sorted(required - names))}.")
    if names - required - optional:
        raise ValueError(f"Columnas desconocidas: {', '.join(sorted(names - required - optional))}.")
    result = {key: value.strip() if isinstance(value, str) else value for key, value in row.items()}
    if kind == "resources":
        data = ResourceCreate.model_validate(result)
        condition = result.get("condition") or "operational"
        if condition not in {"operational", "damaged"}:
            raise ValueError("condition debe ser operational o damaged.")
        return {"id": data.id, "name": data.name.strip(), "category": data.category.strip().lower(),
                "condition": condition}
    space_id = str(result["id"]).upper()
    official = OFFICIAL_SPACES.get(space_id)
    if official is None:
        raise ValueError("El identificador no pertenece al catálogo oficial de salas.")
    raw_capacity = result["capacity"]
    if isinstance(raw_capacity, bool) or not (
        isinstance(raw_capacity, int) or (isinstance(raw_capacity, str) and raw_capacity.isdecimal())
    ):
        raise ValueError("capacity debe ser un entero.")
    capacity = int(raw_capacity)
    if (result["name"], capacity) != (official["name"], official["capacity"]):
        raise ValueError("Nombre o aforo distinto del catálogo oficial; no se sobrescribe.")
    return {"id": space_id, "name": official["name"], "capacity": capacity}


def import_dataset(path: Path, kind: str, db: Session, *, dry_run: bool = False) -> dict[str, Any]:
    """Import valid rows once, report conflicts, and never overwrite existing records."""
    if kind not in FIELDS:
        raise ValueError("Tipo no admitido. Usa resources o spaces.")
    source = path.resolve(strict=True)
    rows = _read_rows(source)
    model = Resource if kind == "resources" else Space
    seen: set[str] = set()
    results: list[dict[str, Any]] = []
    staged: list[Any] = []
    for number, row in rows:
        record_id = str(row.get("id", "")) if isinstance(row, dict) else ""
        try:
            values = _clean(row, kind)
            record_id = values["id"]
            if record_id in seen:
                raise ValueError("Identificador duplicado dentro del archivo.")
            seen.add(record_id)
            existing = db.get(model, record_id)
            if existing is not None:
                if any(getattr(existing, field) != value for field, value in values.items()):
                    raise ValueError("El identificador ya existe con otros datos; no se sobrescribe.")
                status = "skipped"
            else:
                staged.append(model(**values))
                status = "would_import" if dry_run else "imported"
            results.append({"row": number, "id": record_id, "status": status})
        except (ValueError, PydanticValidationError) as exc:
            results.append({"row": number, "id": record_id, "status": "rejected", "reason": str(exc)})
    try:
        if dry_run:
            db.rollback()
        else:
            db.add_all(staged)
            db.commit()
    except Exception:
        db.rollback()
        raise
    return {
        "source": source.name,
        "kind": kind,
        "dry_run": dry_run,
        "summary": {status: sum(item["status"] == status for item in results)
                    for status in ("imported", "would_import", "skipped", "rejected")},
        "rows": results,
    }
